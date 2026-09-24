from openpyxl import Workbook, load_workbook
import os

from flask import Flask, render_template, request
import joblib


# ============================================================
# SAVE DETAILS TO EXCEL
# ============================================================

def save_to_excel(name, phone, income, loan_amount, loan_time, area, prediction):

    file_name = "loan_details.xlsx"

    # Create Excel file if it does not exist
    if not os.path.exists(file_name):

        workbook = Workbook()
        sheet = workbook.active

        sheet.append([
            "Name",
            "Phone Number",
            "Income",
            "Loan Amount",
            "Loan Time",
            "Area",
            "Prediction"
        ])

        workbook.save(file_name)

    # Open Excel file
    workbook = load_workbook(file_name)
    sheet = workbook.active

    # Add submitted details
    sheet.append([
        name,
        phone,
        income,
        loan_amount,
        loan_time,
        area,
        prediction
    ])

    # Make columns wider
    sheet.column_dimensions['A'].width = 20
    sheet.column_dimensions['B'].width = 20
    sheet.column_dimensions['C'].width = 15
    sheet.column_dimensions['D'].width = 20
    sheet.column_dimensions['E'].width = 15
    sheet.column_dimensions['F'].width = 18
    sheet.column_dimensions['G'].width = 25


    # Save Excel file
    workbook.save(file_name)


# ============================================================
# FLASK APP
# ============================================================

app = Flask(__name__)


# Load trained model
model = joblib.load("model.pkl")


# ============================================================
# HOME PAGE
# ============================================================

@app.route('/')
def home():

    return render_template('index.html')


# ============================================================
# PREDICTION PAGE
# ============================================================

@app.route('/predict')
def predict():

    return render_template('predict.html')


# ============================================================
# RESULT PAGE
# ============================================================

@app.route('/result', methods=['POST'])
def result():

    # Get details from form
    name = request.form['applicant_name']

    phone = request.form['applicant_phnnumber']

    income = float(request.form['income'])

    loan_amount = float(request.form['loan_amount'])

    loan_time = int(request.form['loan_time'])

    area = request.form['area']


    # ========================================================
    # CONVERT AREA FOR MACHINE LEARNING MODEL
    # ========================================================

    if area == "Rural":

        area_number = 0

    elif area == "Semiurban":

        area_number = 1

    else:

        area_number = 2


    # ========================================================
    # MAKE PREDICTION
    # ========================================================

    prediction = model.predict(
        [[income, loan_amount, loan_time, area_number]]
    )


    # Convert prediction into readable text
    if prediction[0] == 1:

        result_text = "Loan Approved"

    else:

        result_text = "Loan Rejected"


    # ========================================================
    # SAVE DETAILS TO EXCEL
    # ========================================================

    save_to_excel(
        name,
        phone,
        income,
        loan_amount,
        loan_time,
        area,
        result_text
    )


    # ========================================================
    # SHOW RESULT PAGE
    # ========================================================

    return render_template(
        'result.html',

        name=name,

        phone=phone,

        income=income,

        loan_amount=loan_amount,

        loan_time=loan_time,

        area=area,

        prediction=result_text
    )


# ============================================================
# RUN FLASK
# ============================================================

if __name__ == "__main__":

    app.run(debug=True)